"""Build Crowdfunding_Report.docx from the results/ and figures/ produced by crowdfunding_analysis.ipynb."""
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

RES, FIG = "results", "figures"
S = json.load(open(f"{RES}/summary.json"))
cv = pd.read_csv(f"{RES}/cv_comparison.csv")
test = pd.read_csv(f"{RES}/test_metrics.csv")
coef = pd.read_csv(f"{RES}/ols_coefficients.csv", index_col=0)
stdc = pd.read_csv(f"{RES}/standardised_coefficients.csv", index_col=0)
imp = pd.read_csv(f"{RES}/permutation_importance.csv", index_col=0)
leak = pd.read_csv(f"{RES}/duplicate_leakage.csv", index_col=0)
tgt = pd.read_csv(f"{RES}/target_transform_comparison.csv", index_col=0)

lasso_cv = cv.set_index("model").loc["Lasso"]
dummy_cv = cv.set_index("model").loc["Mean (dummy)"]
lasso_te = test.set_index("model").loc["Lasso"]
bp = S["best_params"]

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)


def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def bullets(items):
    for lead, text in items:
        p = doc.add_paragraph(style="List Bullet")
        if lead:
            p.add_run(lead).bold = True
        p.add_run(text)


def table(df, fmt="{:.3f}", index=True):
    df = df.reset_index() if index else df
    t = doc.add_table(rows=1, cols=len(df.columns))
    t.style = "Light Grid Accent 1"
    for c, name in enumerate(df.columns):
        t.rows[0].cells[c].text = str(name)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for c, v in enumerate(row):
            cells[c].text = fmt.format(v) if isinstance(v, float) else str(v)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def figure(path, caption, width=6.3):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].italic = True
    c.runs[0].font.size = Pt(9)


# ---------------------------------------------------------------- title
doc.add_heading("What Drives Success in Crowdfunding Campaigns?", 0)
para("FINA 4930 — Mini Project 1")

doc.add_heading("Executive summary", 1)
bullets([
    ("Effective sample size is 224, not 439. ",
     f"{S['n_rows_raw'] - S['n']} of the {S['n_rows_raw']} rows are exact duplicates of other campaigns. If the duplicates "
     "are kept, tree models appear to explain about half of the variance, but that is leakage: they memorise copies of "
     "campaigns that sit in both the training and validation folds."),
    ("Selected model: Lasso regression on log(1 + fraction raised). ",
     f"It has the lowest nested-CV error (RMSE {lasso_cv.RMSE:.3f} vs {dummy_cv.RMSE:.3f} for a mean forecast) and is the "
     "sparsest model. Random forests and gradient boosting add no accuracy at this sample size."),
    ("Predictive power is modest. ",
     f"Out-of-fold R² ≈ {lasso_cv.R2_pooled:.2f} and hold-out R² = {lasso_te['R2 (log)']:.2f} (95% CI "
     f"{S['test_r2_ci'][0]:.2f} to {S['test_r2_ci'][1]:.2f}). Observable campaign characteristics explain only a small part "
     "of why campaigns succeed."),
    ("Consistent drivers: ",
     "a pre-order perk (+42% in 1 + fraction raised), an IoT/connected-device product (+71%) and a California base (+34%). "
     "Larger goals are associated with somewhat lower funding ratios. Small-business, generic tech, all-or-nothing "
     "funding and start year are not significant."),
])

# ---------------------------------------------------------------- data
doc.add_heading("1. Data and preparation", 1)
para("The data describe US crowdfunding campaigns from 2011–2013. The dependent variable is frac_raised, the amount raised "
     "divided by the funding goal. Candidate predictors are: pre-order perk, goal amount, technology campaign, small-business "
     "campaign, fixed (all-or-nothing) funding, IoT product, state and start year.")
para("", "Duplicates. ")
doc.paragraphs[-1].add_run(
    f"The campaign identifier (cid) repeats. The {S['n_rows_raw']} rows contain only {S['n']} unique campaigns; one campaign "
    "appears 11 times. Every copy is identical on all columns, so the copies are data-entry duplicates, not separate "
    "funding rounds. Keeping them would weight some campaigns up to 11 times more than others, understate standard "
    "errors, and contaminate cross-validation. Table 1 shows the contamination: under ordinary shuffled CV, the tree "
    "ensembles reach R² ≈ 0.5, but once all copies of a campaign are kept in the same fold their R² falls below zero. "
    "All further analysis uses the de-duplicated sample.")
table(leak.rename(columns=lambda c: c.replace(" (439)", "")))
para("", "Exclusions. ")
doc.paragraphs[-1].add_run(
    "cid is an identifier. delivery (whether the product was eventually delivered) is observed only for a minority of "
    "campaigns and is determined after the funding outcome, so using it as a predictor would be look-ahead leakage.")
para("", "Transformations. ")
doc.paragraphs[-1].add_run(
    "The goal amount is logged. State is reduced to California, New York and Other, because most states have fewer than "
    "10 campaigns. Start year is coded as 2012 and 2013 dummies. The few missing binary values are imputed with the mode "
    "inside each model pipeline. The target is modelled as log(1 + frac_raised). The raw ratio is extremely right-skewed: "
    "23% of campaigns raise nothing and the maximum is 26 times the goal. Table 2 confirms that the log1p model predicts "
    "even the raw ratio better than OLS fitted on the raw ratio.")
table(tgt)
figure(f"{FIG}/01_distributions.png", "Figure 1. Distribution of the target before and after transformation, and of log(goal).")

# ---------------------------------------------------------------- models
doc.add_heading("2. Model candidates and the sample-size constraint", 1)
para("With 179 training campaigns, every extra parameter has a real cost in variance. The candidate set covers the "
     "bias–variance spectrum:")
bullets([
    ("Mean forecast — ", "the benchmark."),
    ("OLS (11 parameters, ~16 observations per parameter) — ", "interpretable; used with HC3 robust errors for inference."),
    ("OLS with all pairwise interactions (56 parameters) — ", "tests for interaction effects, at high over-fitting risk."),
    ("Ridge, Lasso, Elastic Net — ", "shrinkage stabilises estimates for correlated dummies (IoT products are a subset of tech) "
     "and, for the Lasso, also selects variables."),
    ("Huber regression — ", "a robust loss that down-weights the few campaigns raising 10–26 times their goal."),
    ("Random forest and gradient boosting — ", "non-parametric benchmarks that capture non-linearities and interactions."),
])
para("Evaluation design. 20% of campaigns (45) are held out for a single final check. Because a 45-campaign test set is "
     "noisy, model choice rests on nested cross-validation within the training set: hyper-parameters are tuned by 5-fold CV "
     "in an inner loop, and each tuned model is scored in an outer loop of 10-fold CV repeated 3 times (30 folds). The "
     "outer folds never take part in tuning, so the comparison is unbiased.")

doc.add_heading("3. Hyper-parameter selection and tuning", 1)
para("Grids were chosen to bracket sensible values, deliberately favouring simple settings for the tree models. "
     "Values selected by 5-fold CV on the full training set:")
hp = pd.DataFrame([(m, ", ".join(f"{k} = {v:.4g}" if isinstance(v, float) else f"{k} = {v}" for k, v in p.items()))
                   for m, p in bp.items()], columns=["Model", "Selected hyper-parameters"])
table(hp, index=False)
bullets([
    ("Ridge / Lasso: ", "α was searched on a log grid. The validation curves (Figure 2) are flat for small α and turn up "
     "sharply once shrinkage removes real signal. The CV minimum sits just before that point. At the chosen α the Lasso "
     "keeps three predictors: pre-order perk, IoT and California."),
    ("Elastic Net: ", f"the selected l1_ratio of {bp['Elastic Net']['l1_ratio']} is Lasso-like, consistent with a sparse signal."),
    ("Huber: ", "the largest ε and α in the grid were chosen. That means little robustification and strong shrinkage, a sign "
     "that the outliers are not what limits accuracy."),
    ("Random forest: ", "min_samples_leaf is the key regulariser. Validation error falls steeply from 1 to 5–10 observations "
     "per leaf, so fully grown trees over-fit badly."),
    ("Gradient boosting: ", "CV chose stumps (depth 1) with a small learning rate. With depth 2 and lr = 0.05, validation "
     "error rises steadily with more trees (Figure 2). Any interaction structure the boosting adds is noise here."),
])
figure(f"{FIG}/06_validation_curves.png", "Figure 2. Validation curves (5-fold CV RMSE on log(1+frac); band = ±1 sd).")

# ---------------------------------------------------------------- results
doc.add_heading("4. Model outcomes: predictive power and accuracy", 1)
para("Table 4 reports nested-CV results. Pooled R² = 1 − mean(MSE)/Var(y). The average of per-fold R² is also shown, but "
     "with ~18 campaigns per fold it is biased downward.")
table(cv[["model", "RMSE", "RMSE_se", "MAE", "R2_pooled", "R2_mean_fold"]].rename(
    columns={"RMSE_se": "RMSE s.e.", "R2_pooled": "R² pooled", "R2_mean_fold": "R² mean fold"}), index=False)
figure(f"{FIG}/05_cv_comparison.png", "Figure 3. Outer-fold RMSE across the 30 nested-CV folds.", width=5.5)
para("", "Selection (one-standard-error rule). ")
doc.paragraphs[-1].add_run(
    f"The Lasso has the lowest CV RMSE ({lasso_cv.RMSE:.3f}). Every model except the interaction model lies within one "
    f"standard error ({lasso_cv.RMSE_se:.3f}) of it. Under the 1-SE rule we pick the simplest model in that band. The Lasso "
    "is both the most accurate and the sparsest (three active predictors), so it is selected. The trees' lack of advantage "
    "is itself informative: at this sample size there is no detectable non-linear or interaction structure for them to exploit.")
para("", "Does more data help? ")
doc.paragraphs[-1].add_run(
    "The learning curves (Figure 4) show training and validation error converging at about 0.6–0.65 by roughly 100 "
    "campaigns, and the gap for OLS closes as n grows. More campaigns of the same kind would mostly narrow the confidence "
    "intervals. Better prediction would require richer features.")
figure(f"{FIG}/07_learning_curves.png", "Figure 4. Learning curves (repeated 5-fold CV).")
para("Hold-out test set (45 campaigns). Metrics on the original fraction-raised scale use Duan's smearing back-transformation.")
table(test[["model", "R2 (log)", "RMSE (log)", "MAE (frac)", "Median AE (frac)", "Goal hit-rate", "Train R2 (log)"]],
      index=False)
para(f"The Lasso's test R² is {lasso_te['R2 (log)']:.3f}, with a bootstrap 95% CI of [{S['test_r2_ci'][0]:.2f}, "
     f"{S['test_r2_ci'][1]:.2f}]. That interval is wide enough to include zero, which illustrates how little a single small "
     "split can tell us. The typical (median) absolute error is about one times the goal. As a classifier of whether a "
     "campaign reaches its goal, no regression model beats the naive rule \"every campaign succeeds\" "
     f"({S['baseline_hit_rate']:.0%} accuracy). In short, the models identify systematic drivers but cannot forecast "
     "individual campaigns reliably.")
figure(f"{FIG}/08_pred_vs_actual.png", "Figure 5. Lasso: predicted vs. actual and residuals on the test set.", width=5.5)

# ---------------------------------------------------------------- variables
doc.add_heading("5. Relevance of the variables", 1)
para(f"Table 6 reports OLS on all {S['n']} campaigns with HC3 robust standard errors (R² = {S['ols_r2']:.3f}, adjusted R² = "
     f"{S['ols_adj_r2']:.3f}; the regression is jointly significant, p < 0.001). The % effect is exp(b) − 1, the change in "
     "(1 + fraction raised). Variance inflation factors are all below 6. The largest belong to the year dummies, because "
     "few campaigns started in 2011.")
ct = coef[["coef", "robust_se", "p_value", "pct_effect_on_1+frac", "VIF"]].rename(
    columns={"robust_se": "HC3 s.e.", "p_value": "p", "pct_effect_on_1+frac": "% effect"})
table(ct)
para("Standardised coefficients across the linear models and permutation importance on the test set "
     "(the drop in R² when a feature is shuffled):")
table(stdc.join(imp.add_prefix("PI: ")))
figure(f"{FIG}/10_permutation_importance.png", "Figure 6. Permutation importance on the hold-out set.", width=5.5)
bullets([
    ("Pre-order perk (strongest, most robust). ",
     "Campaigns that promise the product are associated with about 42% higher (1 + fraction raised) (p < 0.001). They are "
     "half as likely to raise nothing (16% vs 29%) and more likely to reach their goal (67% vs 55%). Every model ranks it "
     "first or second."),
    ("IoT / connected device. ",
     "+71% (p = 0.012). Median IoT campaigns raise nearly twice their goal, and 81% succeed. This is the most important "
     "feature out of sample. Once IoT is controlled for, the generic tech flag has no effect: the 'tech premium' is really "
     "a connected-device premium."),
    ("California. ",
     "+34% (p = 0.011), consistent with proximity to tech-savvy backer communities and media. The Lasso keeps it, but its "
     "out-of-sample importance is close to zero, so treat it as weaker evidence."),
    ("Goal size. ",
     "Doubling the goal lowers (1 + fraction raised) by about 4.5% (p = 0.06). This is a small, marginal effect, and the "
     "Lasso drops it. Ambitious goals are not heavily penalised in ratio terms."),
    ("Not relevant: ",
     "small-business campaigns, all-or-nothing (fixed) funding, New York and start year. None is significant, and the "
     "Lasso zeroes all of them."),
])
para("", "Delivery (descriptive only). ")
doc.paragraphs[-1].add_run(
    "Among the 53 campaigns with a known delivery status, the 9 that never delivered had raised a median of about 9 times "
    "their goal. Delivered projects raised a median of about 2.2 times. The subsample is very small, but it points to a "
    "risk of over-funding: a campaign that raises far more than planned may struggle to scale up production.")

# ---------------------------------------------------------------- policy
doc.add_heading("6. Policy implications", 1)
doc.add_heading("Crowdfunding portals", 2)
bullets([
    ("Promote reward-based pre-order structures. ",
     "Templates and guidance that help creators offer the product itself as a perk are the most reliable lever in the data."),
    ("Curate and recommend by product type, not broad category. ",
     "Connected-device campaigns perform far better than 'tech' in general. Discovery algorithms should use finer "
     "categories than the tech / small-business split."),
    ("Do not over-promise predictability. ",
     "Observable characteristics explain only around 10–20% of the variation in outcomes, so a portal's 'success score' "
     "built on such fields would mislead creators and backers. Richer signals such as creator track record, prototypes "
     "and early momentum are needed."),
    ("Monitor and manage over-funding. ",
     "Heavily over-funded campaigns appear more likely to fail to deliver. Portals could offer stretch-goal planning tools, "
     "staged disbursement, or fulfilment check-ins once a campaign passes a multiple of its goal."),
    ("Data quality. ",
     "Duplicated campaign records halve the effective sample. Portals that publish data for research or rankings should "
     "de-duplicate it, or analyses will be overconfident."),
])
doc.add_heading("Regulators", 2)
bullets([
    ("Delivery risk is the main consumer-protection issue. ",
     "Backers effectively pre-buy products. Disclosure requirements about production readiness, and about how funds "
     "above the goal will be used, matter more than rules on the funding mechanism itself."),
    ("All-or-nothing funding does not reduce funding success. ",
     "Fixed-funding campaigns raise about the same as flexible ones, so requiring or encouraging all-or-nothing "
     "structures, which protect backers from under-capitalised projects, would not obviously hurt entrepreneurs."),
    ("Avoid policy built on spurious precision. ",
     "Rules calibrated on duplicated or small datasets risk being based on effects that are not there. Regulators should "
     "require standardised, de-duplicated reporting of outcomes, including delivery."),
    ("Geographic concentration. ",
     "The California advantage suggests network effects. Policies that broaden access, such as regional platforms or "
     "investor-education programmes, may support entrepreneurs outside the main hubs."),
])
doc.add_heading("Crowdfunding companies and entrepreneurs", 2)
bullets([
    ("Offer the product as a pre-order perk. ", "This is the clearest controllable driver of higher funding ratios."),
    ("Connected or IoT products resonate with crowdfunding audiences. ",
     "Positioning a product's 'smart' or connected features may matter more than being 'tech' in general."),
    ("Set goals at the true minimum viable amount. ",
     "Larger goals reduce the funding ratio only slightly, but under all-or-nothing rules even a small shortfall means "
     "receiving nothing. Plan production capacity for over-subscription."),
    ("The funding mechanism is not the deciding factor. ",
     "Choose fixed versus flexible funding for its risk-sharing properties, not in the hope of raising more."),
    ("Most of the outcome is driven by what this data cannot see. ",
     "Pitch quality, video, community building and marketing likely dominate, so invest there."),
])

doc.add_heading("7. Limitations", 1)
bullets([
    (None, f"Small sample ({S['n']} unique campaigns, a 45-campaign test set). Effect sizes are imprecise and the CIs are wide."),
    (None, "Observational data: the coefficients are associations. Pre-order perks may proxy for more mature products."),
    (None, "US campaigns from 2011–2013 only. Platform norms have changed since then."),
    (None, "Key quality signals are unobserved. Delivery data are too sparse and too post-hoc to use as a predictor."),
])

doc.save("Crowdfunding_Report.docx")
print("Wrote Crowdfunding_Report.docx")
